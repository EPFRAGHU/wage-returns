import calendar
import difflib
import io
import os
import re
import secrets
from datetime import date, datetime
from functools import wraps

from flask import (Flask, abort, flash, g, jsonify, redirect, render_template,
                   request, send_file, session, url_for)
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import check_password_hash, generate_password_hash

import calc

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-me-in-production")
db_url = os.environ.get("DATABASE_URL", "sqlite:///" + os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "wages.db"))
if db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)
app.config["SQLALCHEMY_DATABASE_URI"] = db_url
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True}
os.makedirs(os.path.join(os.path.dirname(os.path.abspath(__file__)), "data"), exist_ok=True)
db = SQLAlchemy(app)

# ESIC monthly-contribution reason codes for zero working days (from the ESIC MC template)
ESIC_REASONS = [(0, "Without reason"), (1, "On leave"), (2, "Left service"), (3, "Retired"), (4, "Out of coverage"),
                (5, "Expired"), (6, "Non-implemented area"), (7, "Compliance by immediate employer"),
                (8, "Suspension of work"), (9, "Strike / lockout"), (10, "Retrenchment"), (11, "No work"),
                (12, "Doesn't belong to this employer"), (13, "Duplicate IP")]
ESIC_LWD_CODES = {2, 3, 4, 5, 6, 10}  # these need the last working day
ESIC_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "esic_template", "MC_Template11.xls")

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# ───────────────────────── models ─────────────────────────
class Employer(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    est_code = db.Column(db.String(30))
    esic_code = db.Column(db.String(30))
    created = db.Column(db.DateTime, default=datetime.utcnow)


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    pw = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(10), nullable=False)  # admin | employer
    employer_id = db.Column(db.Integer, db.ForeignKey("employer.id"))


class Employee(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    employer_id = db.Column(db.Integer, db.ForeignKey("employer.id"), nullable=False, index=True)
    name = db.Column(db.String(120), nullable=False)
    uan = db.Column(db.String(12))
    ip_no = db.Column(db.String(17))
    dob = db.Column(db.Date)
    doj = db.Column(db.Date)
    actual_basic = db.Column(db.Integer, default=0)
    hra = db.Column(db.Integer, default=0)
    laundry = db.Column(db.Integer, default=0)
    active = db.Column(db.Boolean, default=True)
    sort = db.Column(db.Integer, default=0)


class WageMonth(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    employer_id = db.Column(db.Integer, db.ForeignKey("employer.id"), nullable=False, index=True)
    year = db.Column(db.Integer, nullable=False)
    month = db.Column(db.Integer, nullable=False)
    status = db.Column(db.String(12), default="draft")  # draft | submitted | generated
    submitted_at = db.Column(db.DateTime)
    generated_at = db.Column(db.DateTime)
    __table_args__ = (db.UniqueConstraint("employer_id", "year", "month"),)

    @property
    def label(self):
        return f"{MONTHS[self.month - 1]}-{self.year}"

    @property
    def ym(self):
        return f"{self.year}-{self.month:02d}"


class WageEntry(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    wage_month_id = db.Column(db.Integer, db.ForeignKey("wage_month.id"), nullable=False, index=True)
    employee_id = db.Column(db.Integer, db.ForeignKey("employee.id"), nullable=False)
    actual_basic = db.Column(db.Integer, default=0)
    total_days = db.Column(db.Integer, default=30)
    lop_days = db.Column(db.Integer, default=0)
    hra = db.Column(db.Integer, default=0)
    laundry = db.Column(db.Integer, default=0)
    entered = db.Column(db.Boolean, default=False)  # employer has confirmed this row for the month
    esic_reason = db.Column(db.Integer, default=0)  # ESIC reason code, used only when paid days are 0
    esic_lwd = db.Column(db.Date)                    # ESIC last working day (codes 2,3,4,5,6,10)
    employee = db.relationship("Employee")


class Ceiling(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    effective_ym = db.Column(db.String(10), nullable=False)  # date it takes effect, YYYY-MM-DD
    amount = db.Column(db.Integer, nullable=False)


# ───────────────────────── helpers ─────────────────────────
def _norm(eff):
    """Older rows were stored as YYYY-MM (whole wage month); treat them as the 1st."""
    return eff if len(eff) == 10 else eff + "-01"


def ceiling_detail(year, month):
    """Ceiling for a wage month, pro-rated by calendar days when it changes mid-month.

    Sep-2026: ₹15,000 for 1–16 Sep and ₹25,000 from 17 Sep (S.O. 5109(E) dated
    17.09.2026) → (15,000×16 + 25,000×14) ÷ 30 = ₹19,667.
    Returns (ceiling, parts) where parts = [(from_day, to_day, amount), ...].
    """
    rules = sorted(((_norm(c.effective_ym), c.amount) for c in Ceiling.query.all()))
    days = calendar.monthrange(year, month)[1]
    parts = []
    for d in range(1, days + 1):
        iso = f"{year}-{month:02d}-{d:02d}"
        amt = 15000
        for eff, a in rules:
            if eff <= iso:
                amt = a
        if parts and parts[-1][2] == amt:
            parts[-1][1] = d
        else:
            parts.append([d, d, amt])
    total = sum((p[1] - p[0] + 1) * p[2] for p in parts)
    return calc.r_half_up(calc.Decimal(total) / days), parts


def ceiling_for(year, month):
    return ceiling_detail(year, month)[0]


def ceiling_note(year, month):
    ceil, parts = ceiling_detail(year, month)
    if len(parts) == 1:
        return f"₹{ceil:,}"
    mon = MONTHS[month - 1]
    split = ", ".join(f"₹{a:,} for {f}–{t} {mon}" for f, t, a in parts)
    return f"₹{ceil:,} pro-rata ({split})"


def parse_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def digits(v, n=None):
    if v in (None, ""):
        return None
    if isinstance(v, float):
        v = int(v)
    s = re.sub(r"\D", "", str(v))
    return s or None


def to_int(v, default=0):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return default


def clean_name(n):
    n = re.sub(r"^\s*(MR|MRS|MS|MISS|SMT|SHRI|SRI)\.?\s*", "", str(n or "").strip(), flags=re.I)
    return " ".join(n.upper().split())


def row_from_entry(e: WageEntry) -> calc.Row:
    emp = e.employee
    return calc.Row(name=emp.name, uan=emp.uan, ip_no=emp.ip_no, dob=emp.dob,
                    actual_basic=e.actual_basic or 0, total_days=e.total_days or 0,
                    lop_days=e.lop_days or 0, hra=e.hra or 0, laundry=e.laundry or 0)


def computed_rows(wm: WageMonth):
    ceil = ceiling_for(wm.year, wm.month)
    entries = (WageEntry.query.filter_by(wage_month_id=wm.id).join(Employee)
               .order_by(Employee.sort, Employee.id).all())
    return entries, [calc.compute(row_from_entry(e), wm.year, wm.month, ceil) for e in entries], ceil


# ───────────────────────── auth ─────────────────────────
@app.before_request
def load_user():
    g.user = db.session.get(User, session["uid"]) if "uid" in session else None
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    if request.method == "POST" and request.endpoint != "login":
        tok = request.form.get("csrf") or request.headers.get("X-CSRF")
        if tok != session["csrf"]:
            abort(400, "Form expired. Reload the page and try again.")


def asset(name):
    """Static URL with the file's modified time, so browsers never run an old copy."""
    path = os.path.join(app.static_folder, name)
    v = int(os.path.getmtime(path)) if os.path.exists(path) else 0
    return url_for("static", filename=name, v=v)


@app.context_processor
def inject():
    return {"me": g.get("user"), "csrf": session.get("csrf"), "MONTHS": MONTHS, "asset": asset}


def login_required(role=None):
    def deco(fn):
        @wraps(fn)
        def wrapper(*a, **kw):
            if not g.user:
                return redirect(url_for("login", next=request.path))
            if role and g.user.role != role:
                abort(403)
            return fn(*a, **kw)
        return wrapper
    return deco


def employer_or_404(eid):
    emp = db.session.get(Employer, eid) or abort(404)
    if g.user.role != "admin" and g.user.employer_id != eid:
        abort(403)
    return emp


def month_or_404(eid, mid):
    wm = db.session.get(WageMonth, mid)
    if not wm or wm.employer_id != eid:
        abort(404)
    return wm


def can_edit(wm):
    return g.user.role == "admin" or wm.status == "draft"


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = User.query.filter_by(username=request.form.get("username", "").strip().lower()).first()
        if u and check_password_hash(u.pw, request.form.get("password", "")):
            session.clear()
            session["uid"] = u.id
            return redirect(request.args.get("next") or url_for("home"))
        flash("Wrong username or password.")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required()
def home():
    if g.user.role == "admin":
        return redirect(url_for("admin_home"))
    return redirect(url_for("months", eid=g.user.employer_id))


@app.route("/password", methods=["GET", "POST"])
@login_required()
def change_password():
    if request.method == "POST":
        if not check_password_hash(g.user.pw, request.form["old"]):
            flash("Current password is wrong.")
        elif len(request.form["new"]) < 6:
            flash("New password needs at least 6 characters.")
        else:
            g.user.pw = generate_password_hash(request.form["new"])
            db.session.commit()
            flash("Password changed.")
            return redirect(url_for("home"))
    return render_template("password.html")


# ───────────────────────── admin ─────────────────────────
@app.route("/admin")
@login_required("admin")
def admin_home():
    employers = Employer.query.order_by(Employer.name).all()
    pending = {}
    for e in employers:
        pending[e.id] = WageMonth.query.filter_by(employer_id=e.id, status="submitted").count()
    return render_template("admin.html", employers=employers, pending=pending)


@app.route("/admin/employer", methods=["POST"])
@login_required("admin")
def admin_add_employer():
    f = request.form
    uname = f["username"].strip().lower()
    if User.query.filter_by(username=uname).first():
        flash(f"Username {uname} is taken.")
        return redirect(url_for("admin_home"))
    e = Employer(name=f["name"].strip(), est_code=f.get("est_code", "").strip().upper(),
                 esic_code=f.get("esic_code", "").strip())
    db.session.add(e)
    db.session.flush()
    db.session.add(User(username=uname, pw=generate_password_hash(f["password"]), role="employer", employer_id=e.id))
    db.session.commit()
    flash(f"{e.name} added. Employer signs in as {uname}.")
    return redirect(url_for("admin_home"))


@app.route("/admin/employer/<int:eid>/reset", methods=["POST"])
@login_required("admin")
def admin_reset_pw(eid):
    u = User.query.filter_by(employer_id=eid, role="employer").first_or_404()
    u.pw = generate_password_hash(request.form["password"])
    db.session.commit()
    flash(f"Password reset for {u.username}.")
    return redirect(url_for("admin_home"))


@app.route("/admin/settings", methods=["GET", "POST"])
@login_required("admin")
def admin_settings():
    if request.method == "POST":
        ym, amt = request.form["effective_ym"], to_int(request.form["amount"])
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", ym) and amt > 0:
            row = Ceiling.query.filter_by(effective_ym=ym).first() or Ceiling(effective_ym=ym)
            row.amount = amt
            db.session.add(row)
            db.session.commit()
            flash("Wage ceiling saved.")
    if request.args.get("delete"):
        Ceiling.query.filter_by(id=int(request.args["delete"])).delete()
        db.session.commit()
        return redirect(url_for("admin_settings"))
    today = date.today()
    return render_template("settings.html", ceilings=sorted(Ceiling.query.all(), key=lambda c: _norm(c.effective_ym)),
                           norm=_norm, preview=[(y, m, ceiling_note(y, m)) for y, m in
                                                [(2026, 8), (2026, 9), (2026, 10)]])


# ───────────────────────── months / timeline ─────────────────────────
def timeline(eid):
    wms = WageMonth.query.filter_by(employer_id=eid).order_by(WageMonth.year, WageMonth.month).all()
    counts = dict(db.session.query(WageEntry.wage_month_id, db.func.count(WageEntry.id))
                  .filter(WageEntry.wage_month_id.in_([w.id for w in wms] or [0]))
                  .group_by(WageEntry.wage_month_id).all())
    by_ym = {(w.year, w.month): w for w in wms}
    today = date.today()
    # window: 11 months before the latest of (today, last created) through that month
    end = max([(today.year, today.month)] + list(by_ym.keys()))
    start = min(list(by_ym.keys()) + [(end[0] - 1, end[1] + 1) if end[1] < 12 else (end[0], 1)])
    out, y, m = [], *start
    while (y, m) <= end:
        w = by_ym.get((y, m))
        out.append({"y": y, "m": m, "label": MONTHS[m - 1], "wm": w,
                    "count": counts.get(w.id, 0) if w else 0})
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    peak = max([t["count"] for t in out] + [1])
    for t in out:
        t["h"] = round(18 + 82 * t["count"] / peak) if t["wm"] else 0
    return out


@app.route("/e/<int:eid>/")
@login_required()
def months(eid):
    er = employer_or_404(eid)
    tl = timeline(eid)
    mid = request.args.get("m", type=int)
    wm = db.session.get(WageMonth, mid) if mid else None
    if wm and wm.employer_id != eid:
        abort(404)
    if not wm and not request.args.get("pick"):
        made = [t["wm"] for t in tl if t["wm"]]
        open_ = [w for w in made if w.status == "draft"]
        wm = (open_ or made)[-1] if made else None  # land on the month still being filled
    ctx = {"er": er, "tl": tl, "wm": wm}
    if wm:
        entries, rows, ceil = computed_rows(wm)
        in_month = {e.employee_id for e in entries}
        ctx.update(entries=entries, rows=rows, ceil=ceil, ceil_note=ceiling_note(wm.year, wm.month), editable=can_edit(wm),
                   totals=totals(rows), challan=calc.challan_estimate(rows),
                   addable=Employee.query.filter_by(employer_id=eid, active=True)
                   .filter(~Employee.id.in_(in_month or [0])).order_by(Employee.name).all(),
                   entries_json=[entry_json(e, wm) for e in entries], esic_reasons=ESIC_REASONS,
                   esic_lwd_codes=sorted(ESIC_LWD_CODES), esic_count=sum(1 for r in rows if r.ip_no),
                   pending=sum(1 for e in entries if not e.entered))
    return render_template("months.html", **ctx)


def totals(rows):
    keys = ["actual_basic", "basic", "conveyance", "hra", "laundry", "gross", "epf_wages", "eps_wages",
            "edli_wages", "epf_ee", "eps_er", "epf_er_diff", "esic_wages", "esic_ee", "esic_er", "net"]
    return {k: sum(getattr(r, k) for r in rows) for k in keys}


@app.route("/e/<int:eid>/months/new", methods=["POST"])
@login_required()
def new_month(eid):
    employer_or_404(eid)
    y, m = int(request.form["y"]), int(request.form["m"])
    if WageMonth.query.filter_by(employer_id=eid, year=y, month=m).first():
        flash("That month already exists.")
        return redirect(url_for("months", eid=eid))
    wm = WageMonth(employer_id=eid, year=y, month=m)
    db.session.add(wm)
    db.session.flush()
    prev = (WageMonth.query.filter_by(employer_id=eid)
            .filter((WageMonth.year * 100 + WageMonth.month) < y * 100 + m)
            .order_by(WageMonth.year.desc(), WageMonth.month.desc()).first())
    prev_e = {e.employee_id: e for e in WageEntry.query.filter_by(wage_month_id=prev.id)} if prev else {}
    days = to_int(request.form.get("days")) or 30
    for emp in Employee.query.filter_by(employer_id=eid, active=True).order_by(Employee.sort, Employee.id):
        p = prev_e.get(emp.id)
        db.session.add(WageEntry(wage_month_id=wm.id, employee_id=emp.id,
                                 actual_basic=p.actual_basic if p else emp.actual_basic,
                                 hra=p.hra if p else emp.hra, laundry=p.laundry if p else emp.laundry,
                                 total_days=(p.total_days if p and not request.form.get("days") else days), lop_days=0))
    db.session.commit()
    flash(f"{wm.label} opened with {len(prev_e) if prev else 'all active'} employees. Enter the loss-of-pay days and save.")
    return redirect(url_for("months", eid=eid, m=wm.id))


@app.route("/e/<int:eid>/m/<int:mid>/save", methods=["POST"])
@login_required()
def save_month(eid, mid):
    employer_or_404(eid)
    wm = month_or_404(eid, mid)
    if not can_edit(wm):
        return jsonify(ok=False, error="This month is submitted. Ask the EPF office to reopen it."), 409
    data = request.get_json(force=True)
    entries = {e.id: e for e in WageEntry.query.filter_by(wage_month_id=wm.id)}
    for r in data.get("rows", []):
        e = entries.get(int(r["id"]))
        if not e:
            continue
        e.actual_basic = max(to_int(r.get("actual_basic")), 0)
        e.total_days = min(max(to_int(r.get("total_days"), 30), 1), 31)
        e.lop_days = min(max(to_int(r.get("lop_days")), 0), e.total_days)
        e.hra = max(to_int(r.get("hra")), 0)
        e.laundry = max(to_int(r.get("laundry")), 0)
        dob = parse_date(r.get("dob"))
        if dob and dob != e.employee.dob:
            e.employee.dob = dob  # date of birth lives on the employee master
    if wm.status == "generated":
        wm.status = "submitted"  # data changed after generation; generate again
    db.session.commit()
    return jsonify(ok=True)


@app.route("/e/<int:eid>/m/<int:mid>/rows", methods=["POST"])
@login_required()
def month_rows(eid, mid):
    employer_or_404(eid)
    wm = month_or_404(eid, mid)
    if not can_edit(wm):
        abort(409)
    if request.form.get("add"):
        emp = db.session.get(Employee, int(request.form["add"]))
        if emp and emp.employer_id == eid:
            db.session.add(WageEntry(wage_month_id=wm.id, employee_id=emp.id, actual_basic=emp.actual_basic,
                                     hra=emp.hra, laundry=emp.laundry,
                                     total_days=calendar.monthrange(wm.year, wm.month)[1]))
    if request.form.get("remove"):
        WageEntry.query.filter_by(id=int(request.form["remove"]), wage_month_id=wm.id).delete()
    db.session.commit()
    return redirect(url_for("months", eid=eid, m=mid))


def entry_json(e: WageEntry, wm: WageMonth):
    emp = e.employee
    return {"id": e.id, "employee_id": emp.id, "name": emp.name, "uan": emp.uan or "", "ip_no": emp.ip_no or "",
            "dob": emp.dob.isoformat() if emp.dob else "", "actual_basic": e.actual_basic or 0,
            "total_days": e.total_days or 0, "lop_days": e.lop_days or 0, "hra": e.hra or 0,
            "laundry": e.laundry or 0, "entered": bool(e.entered),
            "esic_reason": e.esic_reason or 0, "esic_lwd": e.esic_lwd.isoformat() if e.esic_lwd else ""}


@app.route("/e/<int:eid>/m/<int:mid>/entry/<int:entry_id>", methods=["POST"])
@login_required()
def save_entry(eid, mid, entry_id):
    """Save one employee's wages for the month (and their UAN / IP / DOB)."""
    employer_or_404(eid)
    wm = month_or_404(eid, mid)
    if not can_edit(wm):
        return jsonify(ok=False, error="This month is locked. Ask the EPF office to reopen it."), 409
    e = WageEntry.query.filter_by(id=entry_id, wage_month_id=wm.id).first_or_404()
    d = request.get_json(force=True)
    uan, ip = digits(d.get("uan")), digits(d.get("ip_no"))
    if uan and len(uan) != 12:
        return jsonify(ok=False, error="UAN must be 12 digits."), 400
    if ip and len(ip) != 10:
        return jsonify(ok=False, error="ESIC IP number must be 10 digits."), 400
    if uan and Employee.query.filter(Employee.employer_id == eid, Employee.uan == uan, Employee.id != e.employee_id).first():
        return jsonify(ok=False, error=f"UAN {uan} already belongs to another employee."), 400
    e.total_days = min(max(to_int(d.get("total_days"), 30), 1), 31)
    e.lop_days = min(max(to_int(d.get("lop_days")), 0), e.total_days)
    e.actual_basic = max(to_int(d.get("actual_basic")), 0)
    e.hra = max(to_int(d.get("hra")), 0)
    e.laundry = max(to_int(d.get("laundry")), 0)
    if e.lop_days >= e.total_days and ip:  # zero paid days: ESIC needs a reason (and a date for some reasons)
        e.esic_reason = to_int(d.get("esic_reason"))
        if e.esic_reason not in dict(ESIC_REASONS):
            return jsonify(ok=False, error="Choose a valid ESIC reason."), 400
        e.esic_lwd = parse_date(d.get("esic_lwd")) if e.esic_reason in ESIC_LWD_CODES else None
        if e.esic_reason in ESIC_LWD_CODES and not e.esic_lwd:
            return jsonify(ok=False, error="ESIC needs the last working day for this reason."), 400
    else:  # ESIC: reason and last working day only apply when no wages are paid
        e.esic_reason, e.esic_lwd = 0, None
    e.entered = True
    emp = e.employee
    if d.get("name", "").strip():
        emp.name = clean_name(d["name"])
    emp.uan, emp.ip_no, emp.dob = uan, ip, parse_date(d.get("dob"))
    # remember the latest pay structure on the master for next month
    emp.actual_basic, emp.hra, emp.laundry = e.actual_basic, e.hra, e.laundry
    if wm.status == "generated":
        wm.status = "submitted"
    db.session.commit()
    return jsonify(ok=True, entry=entry_json(e, wm))


@app.route("/e/<int:eid>/m/<int:mid>/new-employee", methods=["POST"])
@login_required()
def month_new_employee(eid, mid):
    """Create a new employee on the master and add them to this month."""
    employer_or_404(eid)
    wm = month_or_404(eid, mid)
    if not can_edit(wm):
        return jsonify(ok=False, error="This month is locked."), 409
    d = request.get_json(force=True)
    name = clean_name(d.get("name"))
    if not name:
        return jsonify(ok=False, error="Enter the employee's name."), 400
    uan = digits(d.get("uan"))
    if uan and len(uan) != 12:
        return jsonify(ok=False, error="UAN must be 12 digits."), 400
    if uan and Employee.query.filter_by(employer_id=eid, uan=uan).first():
        return jsonify(ok=False, error=f"UAN {uan} is already on the employee list."), 400
    emp = Employee(employer_id=eid, name=name, uan=uan, ip_no=digits(d.get("ip_no")), dob=parse_date(d.get("dob")),
                   sort=(db.session.query(db.func.max(Employee.sort)).filter_by(employer_id=eid).scalar() or 0) + 1)
    db.session.add(emp)
    db.session.flush()
    days = WageEntry.query.filter_by(wage_month_id=wm.id).with_entities(WageEntry.total_days).first()
    e = WageEntry(wage_month_id=wm.id, employee_id=emp.id, total_days=days[0] if days else 30)
    db.session.add(e)
    db.session.commit()
    return jsonify(ok=True, entry=entry_json(e, wm))


@app.route("/e/<int:eid>/m/<int:mid>/days", methods=["POST"])
@login_required()
def month_days(eid, mid):
    """Set total days for every employee in the month at once."""
    employer_or_404(eid)
    wm = month_or_404(eid, mid)
    if not can_edit(wm):
        return jsonify(ok=False, error="This month is locked."), 409
    days = min(max(to_int(request.get_json(force=True).get("days"), 30), 1), 31)
    for e in WageEntry.query.filter_by(wage_month_id=wm.id):
        e.total_days = days
        e.lop_days = min(e.lop_days or 0, days)
    db.session.commit()
    return jsonify(ok=True, days=days)


@app.route("/e/<int:eid>/m/<int:mid>/status", methods=["POST"])
@login_required()
def month_status(eid, mid):
    employer_or_404(eid)
    wm = month_or_404(eid, mid)
    act = request.form["action"]
    if act == "submit" and wm.status == "draft":
        left = WageEntry.query.filter_by(wage_month_id=wm.id, entered=False).count()
        if left:
            flash(f"{left} employee{'s' if left > 1 else ''} still need wages entered. Save each one, or remove those who didn't work this month.")
            return redirect(url_for("months", eid=eid, m=mid))
        wm.status, wm.submitted_at = "submitted", datetime.utcnow()
        flash(f"{wm.label} sent to the EPF office. It is locked now.")
    elif act == "reopen" and g.user.role == "admin":
        wm.status = "draft"
        flash(f"{wm.label} reopened for the employer to edit.")
    elif act == "delete" and wm.status == "draft":
        WageEntry.query.filter_by(wage_month_id=wm.id).delete()
        db.session.delete(wm)
        db.session.commit()
        flash("Month deleted.")
        return redirect(url_for("months", eid=eid))
    db.session.commit()
    return redirect(url_for("months", eid=eid, m=mid))


# ───────────────────────── downloads ─────────────────────────
@app.route("/e/<int:eid>/m/<int:mid>/ecr")
@login_required("admin")
def download_ecr(eid, mid):
    er = employer_or_404(eid)
    wm = month_or_404(eid, mid)
    _, rows, _ = computed_rows(wm)
    txt = calc.build_ecr(rows)
    wm.status, wm.generated_at = "generated", datetime.utcnow()
    db.session.commit()
    fname = f"ECR_{er.est_code or er.id}_{wm.month:02d}{wm.year}.txt"
    return send_file(io.BytesIO(txt.encode("utf-8")), mimetype="text/plain", as_attachment=True, download_name=fname)


@app.route("/e/<int:eid>/m/<int:mid>/esic")
@login_required()
def download_esic(eid, mid):
    """Fill ESIC's own MC template (MC_Template11.xls) for the month, keeping its sheets and formatting."""
    import xlrd
    import xlwt
    from xlutils.copy import copy as xl_copy
    er = employer_or_404(eid)
    wm = month_or_404(eid, mid)
    entries, rows, _ = computed_rows(wm)

    src = xlrd.open_workbook(ESIC_TEMPLATE, formatting_info=True)
    wb = xl_copy(src)
    ws = wb.get_sheet(0)
    # ESIC asks for every column as Text
    txt = xlwt.easyxf("font: name Arial, height 200; align: horiz left", num_format_str="@")
    r_i = 1
    for e, r in zip(entries, rows):
        if not r.ip_no:
            continue
        zero = r.paid_days == 0
        reason = (e.esic_reason or 0) if zero else 0
        lwd = e.esic_lwd.strftime("%d/%m/%Y") if zero and e.esic_lwd and reason in ESIC_LWD_CODES else ""
        name = re.sub(r"[^A-Z ]", "", clean_name(r.name))  # only alphabets and space
        for c, v in enumerate([r.ip_no, name, str(r.paid_days), str(r.esic_wages if not zero else 0), str(reason), lwd]):
            ws.write(r_i, c, v, txt)
        r_i += 1
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.ms-excel", as_attachment=True,
                     download_name=f"ESIC_MC_{er.esic_code or er.id}_{MONTHS[wm.month - 1]}{wm.year}.xls")


@app.route("/e/<int:eid>/m/<int:mid>/register")
@login_required()
def download_register(eid, mid):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    er = employer_or_404(eid)
    wm = month_or_404(eid, mid)
    _, rows, ceil = computed_rows(wm)
    lab = f'{MONTHS[wm.month - 1].upper()}-{wm.year}'
    wb = Workbook()
    ws = wb.active
    ws.title = "Salary statement"
    ws.append([f'SALARY STATEMENT FOR THE MONTH OF "{lab}" — {er.name}'])
    hdr = ["Sl", "Name", "UAN", "Date of birth", "Actual basic", "Total days", "LOP days", "Paid days", "Basic",
           "Conveyance", "House rent", "Laundry", "Gross", "EPF 12%", "ESIC 0.75%", "Total deduction", "Net payment"]
    ws.append(hdr)
    for i, r in enumerate(rows, 1):
        ws.append([i, r.name, r.uan, r.dob.strftime("%d/%m/%Y") if r.dob else "", r.actual_basic, r.total_days,
                   r.total_days - r.paid_days, r.paid_days, r.basic, r.conveyance, r.hra, r.laundry, r.gross,
                   r.epf_ee, r.esic_ee, r.epf_ee + r.esic_ee, r.net])
    t = totals(rows)
    ws.append(["", "TOTAL", "", "", t["actual_basic"], "", "", "", t["basic"], t["conveyance"], t["hra"], t["laundry"],
               t["gross"], t["epf_ee"], t["esic_ee"], t["epf_ee"] + t["esic_ee"], t["net"]])
    ws2 = wb.create_sheet("EPF (UAN)")
    ws2.append([f'UAN BASED STATEMENT FOR "{lab}" (EPF) — EPS/EDLI wage ceiling {ceiling_note(wm.year, wm.month)}'])
    ws2.append(["Sl", "UAN", "Name", "Gross", "EPF wages", "EPS wages", "EDLI wages", "EE 12%", "EPS 8.33%",
                "ER diff 3.67%", "NCP days", "Age note"])
    for i, r in enumerate([r for r in rows if r.uan], 1):
        ws2.append([i, r.uan, r.name, r.gross, r.epf_wages, r.eps_wages, r.edli_wages, r.epf_ee, r.eps_er,
                    r.epf_er_diff, r.total_days - r.paid_days,
                    {"over58": "58+ no EPS", "turns58": "Turns 58 this month", "nodob": "DOB missing"}.get(r.age_flag, "")])
    ws3 = wb.create_sheet("ESIC (IP)")
    ws3.append([f'IP BASED STATEMENT FOR "{lab}" (ESIC)'])
    ws3.append(["Sl", "Name", "IP number", "Paid days", "ESIC wages", "EE 0.75%", "ER 3.25%"])
    for i, r in enumerate([r for r in rows if r.ip_no], 1):
        ws3.append([i, r.name, r.ip_no, r.paid_days, r.esic_wages, r.esic_ee, r.esic_er])
    for sh in (ws, ws2, ws3):
        sh["A1"].font = Font(bold=True, size=12)
        for c in sh[2]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.column_dimensions["B"].width = 30
        sh.column_dimensions["C"].width = 28
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Wage_register_{wm.month:02d}{wm.year}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ───────────────────────── employees ─────────────────────────
@app.route("/e/<int:eid>/employees", methods=["GET", "POST"])
@login_required()
def employees(eid):
    er = employer_or_404(eid)
    if request.method == "POST":
        f = request.form
        emp = db.session.get(Employee, int(f["id"])) if f.get("id") else Employee(employer_id=eid)
        if emp.employer_id != eid:
            abort(403)
        uan = digits(f.get("uan"))
        if uan and len(uan) != 12:
            flash("UAN must be 12 digits.")
            return redirect(url_for("employees", eid=eid))
        ip = digits(f.get("ip_no"))
        if ip and len(ip) != 10:
            flash("ESIC IP number must be 10 digits.")
            return redirect(url_for("employees", eid=eid))
        emp.name = clean_name(f["name"])
        emp.uan, emp.ip_no = uan, ip
        emp.dob, emp.doj = parse_date(f.get("dob")), parse_date(f.get("doj"))
        emp.actual_basic, emp.hra, emp.laundry = to_int(f.get("actual_basic")), to_int(f.get("hra")), to_int(f.get("laundry"))
        emp.active = f.get("active") == "1"
        if not emp.id:
            emp.sort = (db.session.query(db.func.max(Employee.sort)).filter_by(employer_id=eid).scalar() or 0) + 1
        db.session.add(emp)
        db.session.commit()
        flash(f"{emp.name} saved.")
        return redirect(url_for("employees", eid=eid))
    emps = Employee.query.filter_by(employer_id=eid).order_by(Employee.active.desc(), Employee.sort, Employee.id).all()
    today = date.today()
    ages = {e.id: (today.year - e.dob.year - ((today.month, today.day) < (e.dob.month, e.dob.day))) if e.dob else None for e in emps}
    return render_template("employees.html", er=er, emps=emps, ages=ages)


def upsert_employee(eid, name, uan=None, ip=None, dob=None, doj=None, basic=None, hra=None, laundry=None):
    name = clean_name(name)
    q = Employee.query.filter_by(employer_id=eid)
    emp = (q.filter_by(uan=uan).first() if uan else None) or q.filter_by(name=name).first()
    if not emp:
        emp = Employee(employer_id=eid, name=name,
                       sort=(db.session.query(db.func.max(Employee.sort)).filter_by(employer_id=eid).scalar() or 0) + 1)
        db.session.add(emp)
    emp.name = name
    for attr, val in (("uan", uan), ("ip_no", ip), ("dob", dob), ("doj", doj)):
        if val:
            setattr(emp, attr, val)
    for attr, val in (("actual_basic", basic), ("hra", hra), ("laundry", laundry)):
        if val is not None:
            setattr(emp, attr, val)
    return emp


@app.route("/e/<int:eid>/employees/import", methods=["POST"])
@login_required()
def import_employees(eid):
    from openpyxl import load_workbook
    employer_or_404(eid)
    f = request.files.get("file")
    if not f:
        flash("Choose the employee Excel file first.")
        return redirect(url_for("employees", eid=eid))
    ws = load_workbook(f, data_only=True).active
    header_row, cols = None, {}
    for r in ws.iter_rows(min_row=1, max_row=10):
        names = [str(c.value or "").strip().lower() for c in r]
        if any("name" == n or n.startswith("name") for n in names):
            header_row = r[0].row
            for i, n in enumerate(names):
                for key, pats in {"name": ["name"], "uan": ["uan"], "ip": ["ip", "esic"], "dob": ["birth", "dob"],
                                  "doj": ["joining", "doj"], "basic": ["actual basic", "basic"],
                                  "hra": ["house", "hra"], "laundry": ["laundry", "washing"]}.items():
                    if key not in cols and any(p in n for p in pats):
                        cols[key] = i
            break
    if header_row is None or "name" not in cols:
        flash("Couldn't find a Name column. Use the employee master template.")
        return redirect(url_for("employees", eid=eid))
    n = 0
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        get = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
        nm = get("name")
        if not nm or str(nm).strip().upper() == "TOTAL":
            continue
        upsert_employee(eid, nm, digits(get("uan")), digits(get("ip")), parse_date(get("dob")), parse_date(get("doj")),
                        to_int(get("basic"), None) if get("basic") is not None else None,
                        to_int(get("hra"), None) if get("hra") is not None else None,
                        to_int(get("laundry"), None) if get("laundry") is not None else None)
        n += 1
    db.session.commit()
    flash(f"{n} employees imported or updated.")
    return redirect(url_for("employees", eid=eid))


@app.route("/e/<int:eid>/employees/template")
@login_required()
def employee_template(eid):
    er = employer_or_404(eid)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["Sl", "Name", "UAN", "ESIC IP number", "Date of birth (DD/MM/YYYY)", "Date of joining (DD/MM/YYYY)",
               "Actual basic pay", "House rent", "Laundry allowance"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, e in enumerate(Employee.query.filter_by(employer_id=eid, active=True).order_by(Employee.sort), 1):
        ws.append([i, e.name, e.uan, e.ip_no, e.dob.strftime("%d/%m/%Y") if e.dob else "",
                   e.doj.strftime("%d/%m/%Y") if e.doj else "", e.actual_basic, e.hra, e.laundry])
    for col, w in zip("ABCDEFGHI", (5, 30, 15, 15, 16, 16, 12, 12, 12)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Employee_master_{er.est_code or er.id}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ───────────────────── import the client's monthly salary workbook ─────────────────────
def import_salary_workbook(eid, fileobj):
    """Reads the client's 3-sheet format (salary statement / UAN list / IP list).

    Returns (WageMonth, messages). Employees are matched to UAN/IP rows by name.
    """
    from openpyxl import load_workbook
    wb = load_workbook(fileobj, data_only=False)
    ws1 = wb.worksheets[0]
    title = str(ws1["A1"].value or "")
    mt = re.search(r"([A-Za-z]{3,9})[\s\-']*(\d{4})", title)
    if not mt:
        raise ValueError("Couldn't read the month from cell A1.")
    mon = next(i for i, m in enumerate(MONTHS, 1) if mt.group(1)[:3].lower() == m.lower())
    year = int(mt.group(2))

    uans, ips = {}, {}
    if len(wb.worksheets) > 1:
        for r in wb.worksheets[1].iter_rows(min_row=4, values_only=True):
            if r[1] and r[2] and digits(r[1]) and len(digits(r[1])) == 12:
                uans[clean_name(r[2])] = digits(r[1])
    if len(wb.worksheets) > 2:
        for r in wb.worksheets[2].iter_rows(min_row=4, values_only=True):
            if r[1] and r[2] and digits(r[2]):
                ips[clean_name(r[1])] = digits(r[2])

    def match(name, table):
        if name in table:
            return table[name]
        best = difflib.get_close_matches(name, list(table), n=1, cutoff=0.85)
        return table[best[0]] if best else None

    wm = WageMonth.query.filter_by(employer_id=eid, year=year, month=mon).first()
    if wm:
        WageEntry.query.filter_by(wage_month_id=wm.id).delete()
    else:
        wm = WageMonth(employer_id=eid, year=year, month=mon)
        db.session.add(wm)
        db.session.flush()
    # find columns by their headings so an added column (e.g. DATE OF BIRTH) doesn't break the import
    defaults = {"name": 1, "basic": 2, "days": 3, "lop": 4, "hra": 8, "laundry": 9}
    pats = {"name": "NAME", "basic": "ACTUAL BASIC", "days": "TOTAL DAYS", "lop": "LOP", "hra": "HOUSE",
            "laundry": "LAUNDRY", "dob": "BIRTH", "uan": "UAN", "ip": "IP NO"}
    col, hdr_row = {}, 2
    for hr in ws1.iter_rows(min_row=1, max_row=6):
        texts = [" ".join(str(c.value or "").upper().split()) for c in hr]
        if "NAME" in texts:
            hdr_row = hr[0].row
            for i, t in enumerate(texts):
                for key, pat in pats.items():
                    if key not in col and ((t == pat) if key == "name" else (pat in t)):
                        col[key] = i
            break
    for k, v in defaults.items():
        col.setdefault(k, v)
    n = 0
    for r in ws1.iter_rows(min_row=hdr_row + 1, values_only=True):
        get = lambda k: r[col[k]] if k in col and col[k] < len(r) else None
        name = get("name")
        if not name or not isinstance(name, str) or name.strip().upper() in ("TOTAL", "NAME") or not re.search(r"[A-Za-z]", name):
            continue
        nm = clean_name(name)
        uan = digits(get("uan")) or match(nm, uans)
        ip = digits(get("ip")) or match(nm, ips)
        emp = upsert_employee(eid, nm, uan, ip, dob=parse_date(get("dob")), basic=to_int(get("basic")),
                              hra=to_int(get("hra")), laundry=to_int(get("laundry")))
        db.session.flush()
        db.session.add(WageEntry(wage_month_id=wm.id, employee_id=emp.id, actual_basic=to_int(get("basic")),
                                 total_days=to_int(get("days"), 30), lop_days=to_int(get("lop")), hra=to_int(get("hra")),
                                 laundry=to_int(get("laundry")), entered=True))
        n += 1
    db.session.commit()
    return wm, n


@app.route("/e/<int:eid>/import-month", methods=["POST"])
@login_required()
def import_month(eid):
    employer_or_404(eid)
    f = request.files.get("file")
    try:
        wm, n = import_salary_workbook(eid, f)
        flash(f"{wm.label}: {n} employees imported from the salary sheet. Check dates of birth before generating.")
        return redirect(url_for("months", eid=eid, m=wm.id))
    except Exception as ex:  # noqa: BLE001
        flash(f"Import failed: {ex}")
        return redirect(url_for("months", eid=eid))


# ───────────────────────── bootstrap ─────────────────────────
def init_db():
    db.create_all()
    cols = [c["name"] for c in db.inspect(db.engine).get_columns("wage_entry")]
    for col, ddl in (("entered", "BOOLEAN DEFAULT FALSE"), ("esic_reason", "INTEGER DEFAULT 0"), ("esic_lwd", "DATE")):
        if col not in cols:  # databases created by earlier versions
            with db.engine.begin() as con:
                con.execute(db.text(f"ALTER TABLE wage_entry ADD COLUMN {col} {ddl}"))
    if not User.query.filter_by(role="admin").first():
        db.session.add(User(username=os.environ.get("ADMIN_USER", "admin"),
                            pw=generate_password_hash(os.environ.get("ADMIN_PASSWORD", "admin123")), role="admin"))
    if db.engine.dialect.name == "postgresql":
        with db.engine.begin() as con:
            con.execute(db.text("ALTER TABLE ceiling ALTER COLUMN effective_ym TYPE VARCHAR(10)"))
    if not Ceiling.query.first():
        db.session.add(Ceiling(effective_ym="2014-09-01", amount=15000))
        db.session.add(Ceiling(effective_ym="2026-09-17", amount=25000))
    # earlier versions seeded ₹25,000 from the Oct-2026 wage month; the notification applies from 17.09.2026
    old = Ceiling.query.filter_by(effective_ym="2026-10", amount=25000).first()
    if old and not Ceiling.query.filter_by(effective_ym="2026-09-17").first():
        old.effective_ym = "2026-09-17"
    db.session.commit()


with app.app_context():
    init_db()

if __name__ == "__main__":
    app.run(debug=True, port=5000)
